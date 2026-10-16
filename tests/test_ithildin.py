#!/usr/bin/env python3
"""Tests for ithildin.py. Run: python3 tests/test_ithildin.py

Builds a 126-page maintenance manual in four formats with planted facts (tests/ithildin_fixtures.py)
and checks that map, find, read and table locate each fact in the right section, with exact text.
Needs reportlab and python-docx to build the fixtures; tests that need them are skipped otherwise.
"""
import contextlib
import io
import os
import re
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "ithildin", "scripts"))
sys.path.insert(0, os.path.join(ROOT, "tests"))
import ithildin  # noqa: E402
import ithildin_fixtures as fx  # noqa: E402

TMP = tempfile.mkdtemp(prefix="ithildin-test-")
os.environ["ITHILDIN_CACHE"] = os.path.join(TMP, "cache")
HAVE_BUILDERS = True
try:
    blocks = fx.build_model()
    fx.write_markdown(blocks, os.path.join(TMP, "manual.md"))
    fx.write_docx(blocks, os.path.join(TMP, "manual.docx"))
    fx.write_pdf(blocks, os.path.join(TMP, "manual.pdf"), bookmarks=True)
    fx.write_pdf(blocks, os.path.join(TMP, "manual-flat.pdf"), bookmarks=False)
except ImportError:
    HAVE_BUILDERS = False

FORMATS = ["manual.pdf", "manual-flat.pdf", "manual.docx", "manual.md"]


def run(*argv):
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        ithildin.main(list(argv))
    return out.getvalue()


def path(name):
    return os.path.join(TMP, name)


@unittest.skipUnless(HAVE_BUILDERS, "reportlab and python-docx are needed to build fixtures")
class Structure(unittest.TestCase):
    def test_every_format_finds_the_numbered_sections(self):
        for f in FORMATS:
            with self.subTest(f=f):
                m = run("map", path(f), "--budget", "9000")
                for sid, title in [("1.3", "Definitions"), ("7.3", "Operating Limits"),
                                   ("9.4", "Torque Specifications"), ("11.2", "Exclusions")]:
                    self.assertRegex(m, rf"(?m)^\s*{re.escape(sid)} {title} \|", f)
                self.assertIn("Appendix B Parts List", m)

    def test_flat_pdf_matches_bookmarked_pdf(self):
        def ranges(f):
            return {x.group(1): x.group(2) for x in
                    re.finditer(r"(?m)^\s*(\d+\.\d+) [^|]+\| (p[\d-]+) \|", run("map", path(f), "--budget", "9000"))}
        a, b = ranges("manual.pdf"), ranges("manual-flat.pdf")
        self.assertGreaterEqual(len(a), 40)
        self.assertEqual(a, b)

    def test_outline_source_is_reported(self):
        self.assertIn("PDF bookmarks", run("map", path("manual.pdf")))
        self.assertIn("no bookmarks in the file", run("map", path("manual-flat.pdf")))
        self.assertIn("Word heading styles", run("map", path("manual.docx")))

    def test_running_headers_and_contents_page_are_not_sections(self):
        doc = ithildin.load(path("manual-flat.pdf"))
        self.assertFalse(any("Rev C" in t for _, t in doc.lines))
        nums = [h["num"] for h in doc.headings if h["num"]]
        self.assertEqual(len(nums), len(set(nums)), "a section number appears twice (contents page?)")

    def test_map_is_small_and_respects_budget(self):
        for f in FORMATS:
            with self.subTest(f=f):
                doc = ithildin.load(path(f))
                full = ithildin.est_tokens("\n".join(t for _, t in doc.lines))
                m = run("map", path(f))
                self.assertLess(ithildin.est_tokens(m), 2700)
                self.assertLess(ithildin.est_tokens(m), full / 25)
                small = run("map", path(f), "--budget", "800")
                self.assertLess(ithildin.est_tokens(small), 1100)
                self.assertIn("7 Cooling System", small)

    def test_defined_terms(self):
        m = run("map", path("manual.pdf"))
        self.assertIn("Service Interval (p9)", m)
        self.assertIn("Authorized Dealer", m)


@unittest.skipUnless(HAVE_BUILDERS, "reportlab and python-docx are needed to build fixtures")
class Retrieval(unittest.TestCase):
    def test_find_locates_every_planted_fact_in_its_section(self):
        for f in FORMATS:
            for key, fact in fx.PLANTED.items():
                with self.subTest(f=f, fact=key):
                    out = run("find", path(f), fact["needle"].split(";")[0][:40])
                    self.assertRegex(out, r"\b[1-9]\d* hits", out[:300])
                    if fact["section"][0].isdigit():
                        self.assertIn(f"§{fact['section']} ", out)

    def test_find_batches_patterns_and_reports_counts(self):
        out = run("find", path("manual.pdf"), "coolant outlet", "altitude", "OF-7731")
        self.assertRegex(out.splitlines()[0], r"'coolant outlet': 1 hits.*'altitude': 1 hits.*'OF-7731': 1 hits")

    def test_find_matches_phrases_wrapped_across_lines(self):
        doc = ithildin.load(path("manual.pdf"))
        plain = sum("lightly oiled threads" in t for _, t in doc.lines)
        out = run("find", path("manual.pdf"), "lightly oiled threads")
        found = int(re.search(r"(\d+) hits", out).group(1))
        self.assertGreater(found, plain)

    def test_read_section_returns_exact_text_with_pages(self):
        out = run("read", path("manual.pdf"), "7.3")
        self.assertTrue(out.startswith("[ithildin read: 7.3 Operating Limits, p71-73]"), out[:80])
        self.assertIn("[p72]", out)
        flat = re.sub(r"\s+", " ", out)
        self.assertIn("maximum coolant outlet temperature is 104 °C", flat)
        self.assertNotIn("7.4 Draining", out)

    def test_read_pages_and_lines(self):
        self.assertIn("[p108]", run("read", path("manual.pdf"), "p108"))
        out = run("read", path("manual.docx"), "L5-6")
        self.assertEqual(len(out.strip().splitlines()), 3)

    def test_read_truncates_with_exact_continuation(self):
        out = run("read", path("manual.pdf"), "7", "--max-tokens", "500")
        m = re.search(r"Continue with: read FILE (L\d+-\d+)", out)
        self.assertTrue(m, out[-300:])
        self.assertIn("by subsection: 7.2, 7.3, 7.4", out[-200:])  # remaining subsections offered
        nxt = run("read", path("manual.pdf"), m.group(1), "--max-tokens", "200")
        last = [l for l in out.splitlines() if not l.startswith("[")][-1]
        self.assertNotIn(last, nxt.splitlines()[1:3])

    def test_read_rejects_unknown_targets_helpfully(self):
        with self.assertRaises(SystemExit) as cm:
            run("read", path("manual.pdf"), "99.9")
        self.assertIn("Run map", str(cm.exception))

    def test_table_returns_exact_rows_in_every_format(self):
        for f in FORMATS:
            with self.subTest(f=f):
                tid = re.search(r"(?m)^\s*(T\d+) [^:]+: Table \d+\. Torque values for engine",
                                run("map", path(f), "--budget", "20000")).group(1)
                out = run("table", path(f), tid)
                self.assertIn("| Cylinder head bolt | M12 | 85 | Spiral, 3 passes |", out)

    def test_parts_table(self):
        out = run("find", path("manual.docx"), "Common replacement parts")
        self.assertIn("1 hits", out)


@unittest.skipUnless(HAVE_BUILDERS, "reportlab is needed")
class EdgeCases(unittest.TestCase):
    def test_scanned_pdf_gets_an_ocr_note(self):
        from reportlab.pdfgen import canvas
        p = path("scan.pdf")
        c = canvas.Canvas(p)
        for _ in range(5):
            c.rect(50, 50, 400, 600, fill=1)
            c.showPage()
        c.save()
        self.assertIn("scanned PDF", run("map", p))

    def test_plain_text_without_structure_gives_samples(self):
        p = path("notes.txt")
        with open(p, "w") as f:
            f.write("\n".join(f"meeting note {i}: we discussed item {i} and agreed to follow up." for i in range(400)))
        m = run("map", p)
        self.assertIn("No section structure found", m)
        self.assertIn("L1:", m)

    def test_html_headings(self):
        p = path("page.html")
        with open(p, "w") as f:
            f.write("<html><body><nav>menu</nav><h1>Guide</h1><p>Intro text here.</p><h2>Setup</h2>"
                    "<p>Install the tool first.</p><script>var x=1;</script><h2>Usage</h2><p>Run it.</p></body></html>")
        m = run("map", p)
        self.assertIn("Setup", m)
        self.assertNotIn("var x", run("read", p, "L1-20"))

    def test_pptx_slides_are_sections(self):
        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx not installed")
        prs = Presentation()
        for i in range(1, 6):
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = f"Quarterly result {i}"
            s.placeholders[1].text = f"Revenue grew {i * 3}% in region {i}."
        p = path("deck.pptx")
        prs.save(p)
        m = run("map", p)
        self.assertIn("5 slides", m)
        self.assertIn("Quarterly result 3 | p3", m)
        self.assertIn("p4", run("find", p, "region 4"))

    def test_cache_is_reused_and_keyed_on_content(self):
        p = path("small.md")
        with open(p, "w") as f:
            f.write("# A\n\nOne.\n\n# B\n\nTwo.\n")
        self.assertIn("B | L3-4", run("map", p))
        with open(p, "w") as f:
            f.write("# A\n\nOne.\n\n# C\n\nThree.\n")
        self.assertIn("C | L3-4", run("map", p))


if __name__ == "__main__":
    unittest.main(verbosity=1)
