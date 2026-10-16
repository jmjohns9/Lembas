#!/usr/bin/env python3
"""Build long test documents with planted facts for ithildin.

    python3 tests/ithildin_fixtures.py OUT_DIR

Writes manual.pdf (with bookmarks), manual-flat.pdf (no bookmarks), manual.docx and manual.md:
a ~150-page maintenance manual with a dotted table of contents, running headers and footers,
numbered sections, tables with captions, a definitions section and an appendix. PLANTED lists the
facts and where they live, so tests and benchmarks can check answers against the source.
"""
import os
import random
import sys

PLANTED = {
    "coolant_max": {"answer": "104 °C", "section": "7.3", "needle": "maximum coolant outlet temperature is 104 °C"},
    "head_bolt_torque": {"answer": "85 N·m", "section": "9.4", "needle": "Torque values for engine fasteners"},
    "altitude_void": {"answer": "3,200 m", "section": "11.2", "needle": "operated above 3,200 m"},
    "service_interval": {"answer": "500 operating hours", "section": "1.3", "needle": "means 500 operating hours"},
    "oil_filter_pn": {"answer": "OF-7731-B", "section": "Appendix B", "needle": "OF-7731-B"},
}

SUBJECTS = ["the pump assembly", "the controller", "the drive belt", "the intake manifold", "the sensor harness",
            "the hydraulic line", "the filter housing", "the gearbox", "the relief valve", "the mounting bracket",
            "the coolant loop", "the bearing set", "the exhaust bypass", "the fuel rail", "the control panel"]
VERBS = ["inspect", "verify", "record", "clean", "replace", "tighten", "measure", "lubricate", "calibrate", "check"]
OBJS = ["for wear and scoring", "against the reference values", "before each shift", "after the first 50 hours",
        "using the approved solvent", "with the unit isolated", "to the specified torque", "for leaks at the joints",
        "with a calibrated gauge", "and log the reading in the maintenance record"]
NOTES = ["Operators should allow the unit to cool before opening any panel.",
         "Use only parts listed in the illustrated parts catalog.",
         "Readings outside tolerance must be reported to the site supervisor.",
         "Do not reuse gaskets or self-locking nuts.",
         "Wear eye protection when working near pressurized lines.",
         "Keep the work area free of loose tools and rags.",
         "Torque values apply to clean, lightly oiled threads.",
         "Allow sealant to cure for at least four hours before startup."]

CHAPTERS = [
    ("1", "General Information", ["Scope", "Safety Conventions", "Definitions", "Related Documents"]),
    ("2", "Product Description", ["Overview", "Major Assemblies", "Operating Principles"]),
    ("3", "Installation", ["Site Requirements", "Foundation and Mounting", "Electrical Connections", "Commissioning"]),
    ("4", "Operation", ["Controls and Indicators", "Starting Procedure", "Normal Operation", "Shutdown"]),
    ("5", "Scheduled Maintenance", ["Daily Checks", "Weekly Checks", "Service Interval Tasks", "Annual Overhaul"]),
    ("6", "Lubrication System", ["Oil Specification", "Oil Change Procedure", "Filter Replacement"]),
    ("7", "Cooling System", ["Description", "Coolant Specification", "Operating Limits", "Draining and Refilling"]),
    ("8", "Fuel System", ["Fuel Requirements", "Injector Service", "Fuel Filter Service"]),
    ("9", "Engine Mechanical", ["Cylinder Head Removal", "Valve Adjustment", "Timing Belt", "Torque Specifications"]),
    ("10", "Electrical System", ["Wiring Overview", "Battery Care", "Fault Codes"]),
    ("11", "Warranty", ["Coverage", "Exclusions", "Claims Procedure"]),
    ("12", "Troubleshooting", ["Engine Will Not Start", "Overheating", "Low Oil Pressure", "Excessive Vibration"]),
]


def para(rng, n=5):
    out = []
    for _ in range(n):
        if rng.random() < 0.18:
            out.append(rng.choice(NOTES))
        else:
            s = f"{rng.choice(VERBS).capitalize()} {rng.choice(SUBJECTS)} {rng.choice(OBJS)}"
            if rng.random() < 0.4:
                s += f"; the nominal value is {rng.randint(10, 900)} {rng.choice(['kPa', 'mm', 'rpm', 'N·m', 'L/min', 'V'])}"
            out.append(s + ".")
    return " ".join(out)


def build_model(seed=7, paras_per_section=22):
    """Return a list of blocks: ('h', level, num, title) | ('p', text) | ('table', caption, rows)."""
    rng = random.Random(seed)
    blocks = []
    tcount = 0
    for num, title, subs in CHAPTERS:
        blocks.append(("h", 1, num, title))
        blocks.append(("p", para(rng, 4)))
        for j, sub in enumerate(subs, 1):
            snum = f"{num}.{j}"
            blocks.append(("h", 2, snum, sub))
            for k in range(paras_per_section):
                text = para(rng, 6)
                if snum == "7.3" and k == paras_per_section // 2:
                    text += (" During continuous duty the maximum coolant outlet temperature is 104 °C; "
                             "if the gauge exceeds this value, reduce load and inspect the radiator core.")
                if snum == "11.2" and k == 3:
                    text += (" The warranty is void if the unit is operated above 3,200 m altitude without the "
                             "high-altitude kit, or if non-approved coolant is used.")
                if snum == "1.3" and k == 1:
                    text += (' "Service Interval" means 500 operating hours or twelve months, whichever comes first. '
                             '"Authorized Dealer" means a dealer listed on the manufacturer\'s portal.')
                blocks.append(("p", text))
                if k in (2, 6) and rng.random() < 0.6 or (snum == "9.4" and k == 4):
                    tcount += 1
                    if snum == "9.4" and k == 4:
                        rows = [["Fastener", "Size", "Torque (N·m)", "Sequence"],
                                ["Cylinder head bolt", "M12", "85", "Spiral, 3 passes"],
                                ["Rocker cover screw", "M6", "10", "Cross pattern"],
                                ["Exhaust manifold nut", "M10", "48", "Center out"],
                                ["Oil pan bolt", "M8", "24", "Cross pattern"]]
                        cap = f"Table {tcount}. Torque values for engine fasteners"
                    else:
                        rows = [["Item", "Interval", "Limit", "Action"]] + [
                            [rng.choice(SUBJECTS).replace("the ", "").title(), f"{rng.randint(1, 20) * 50} h",
                             f"{rng.randint(1, 99)} {rng.choice(['mm', 'kPa', '%'])}", rng.choice(VERBS).capitalize()]
                            for _ in range(5)]
                        cap = f"Table {tcount}. {rng.choice(['Inspection limits', 'Service checks', 'Wear limits', 'Reference values'])} for {sub.lower()}"
                    blocks.append(("table", cap, rows))
    blocks.append(("h", 1, "", "Appendix A Tools and Consumables"))
    for _ in range(6):
        blocks.append(("p", para(rng, 6)))
    blocks.append(("h", 1, "", "Appendix B Parts List"))
    blocks.append(("p", "Order parts through an Authorized Dealer. Quote the part number and the unit serial number."))
    rows = [["Part", "Part number", "Qty"], ["Air filter element", "AF-2210-C", "1"], ["Oil filter", "OF-7731-B", "1"],
            ["Fuel filter", "FF-4402-A", "2"], ["Drive belt", "DB-1180-K", "1"], ["Coolant hose kit", "CH-9020-D", "1"]]
    blocks.append(("table", f"Table {tcount + 1}. Common replacement parts", rows))
    for _ in range(4):
        blocks.append(("p", para(rng, 6)))
    return blocks


def write_markdown(blocks, path):
    with open(path, "w") as f:
        f.write("# ACME X-200 Maintenance Manual\n\n")
        for b in blocks:
            if b[0] == "h":
                f.write(f"{'#' * (b[1] + 1)} {(b[2] + ' ') if b[2] else ''}{b[3]}\n\n")
            elif b[0] == "p":
                f.write(b[1] + "\n\n")
            else:
                f.write(b[1] + "\n\n")
                rows = b[2]
                f.write("| " + " | ".join(rows[0]) + " |\n|" + "---|" * len(rows[0]) + "\n")
                for r in rows[1:]:
                    f.write("| " + " | ".join(r) + " |\n")
                f.write("\n")


def write_docx(blocks, path):
    import docx
    d = docx.Document()
    d.add_heading("ACME X-200 Maintenance Manual", 0)
    for b in blocks:
        if b[0] == "h":
            d.add_heading(f"{(b[2] + ' ') if b[2] else ''}{b[3]}", b[1])
        elif b[0] == "p":
            d.add_paragraph(b[1])
        else:
            d.add_paragraph(b[1])
            rows = b[2]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
    d.save(path)


def write_pdf(blocks, path, bookmarks=True):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font = "Helvetica"
    for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",):
        if os.path.exists(cand):
            pdfmetrics.registerFont(TTFont("DejaVu", cand))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", cand.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")))
            font = "DejaVu"
    ss = getSampleStyleSheet()
    body = ss["BodyText"].clone("b", fontName=font, fontSize=10, leading=13)
    h1 = ss["Heading1"].clone("h1", fontName=font + ("-Bold" if font == "DejaVu" else "-Bold"))
    h2 = ss["Heading2"].clone("h2", fontName=font + ("-Bold" if font == "DejaVu" else "-Bold"))

    class Doc(BaseDocTemplate):
        def __init__(self, fn, **kw):
            super().__init__(fn, pagesize=letter, **kw)
            frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="f")
            self.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=self.decorate)])
            self.n = 0

        def decorate(self, canvas, doc):
            canvas.saveState()
            canvas.setFont(font, 8)
            canvas.drawString(72, letter[1] - 40, "ACME X-200 Maintenance Manual  Rev C")
            canvas.drawRightString(letter[0] - 72, 30, f"Page {doc.page}")
            canvas.restoreState()

        def afterFlowable(self, fl):
            if isinstance(fl, Paragraph) and fl.style.name in ("h1", "h2"):
                text = fl.getPlainText()
                level = 0 if fl.style.name == "h1" else 1
                self.notify("TOCEntry", (level, text, self.page))
                if bookmarks:
                    key = f"k{self.n}"
                    self.n += 1
                    self.canv.bookmarkPage(key)
                    self.canv.addOutlineEntry(text, key, level=level, closed=False)

    story = [Paragraph("ACME X-200 Maintenance Manual", ss["Title"].clone("t", fontName=font)),
             Paragraph("Contents", ss["Heading2"].clone("c", fontName=font))]
    toc = TableOfContents(dotsMinLevel=0)
    toc.levelStyles = [body.clone("toc1"), body.clone("toc2", leftIndent=16)]
    story += [toc, PageBreak()]
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(f"{(b[2] + ' ') if b[2] else ''}{b[3]}", h1 if b[1] == 1 else h2))
        elif b[0] == "p":
            story.append(Paragraph(b[1], body))
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(b[1], body))
            t = Table(b[2], hAlign="LEFT")
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
            story += [t, Spacer(1, 8)]
    Doc(path).multiBuild(story)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "fixtures"
    os.makedirs(out, exist_ok=True)
    blocks = build_model()
    write_markdown(blocks, os.path.join(out, "manual.md"))
    write_docx(blocks, os.path.join(out, "manual.docx"))
    write_pdf(blocks, os.path.join(out, "manual.pdf"), bookmarks=True)
    write_pdf(blocks, os.path.join(out, "manual-flat.pdf"), bookmarks=False)
    print("wrote", out)


if __name__ == "__main__":
    main()
